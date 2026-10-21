"""Review from a sheet: `politricks review <bill-id> --from-sheet <file.xlsx>`.

Why this exists: the learner reviews a whole bill in Excel (the sheet that
`politricks export` writes) rather than one item at a time in the terminal.
He decided (2026-10-15) that a row he marks **Accept** in that sheet IS his
accept, as long as he is the one running this command. So this is a second
way into the same human gate as the walk, not a way around it:

- It only runs from `politricks review`, at a real terminal: with input piped
  in (a script, an agent, `echo yes |`) it refuses before reading anything.
  It saves nothing until he types `yes` at the end; anything else, end of
  input or Ctrl-C quits with the bill untouched.
- It accepts only what he saw. Every row carries the fingerprint the item had
  at export; a row whose item changed in the JSON since is skipped. A group
  whose section list moves because of tag edits in the same sheet can't be
  accepted (he read the old list). The plan shows every new text in full
  before he says yes.
- It refuses a sheet that may not say what it seems to: a column sorted on
  its own (a row's legal text isn't its section's, or an edited text is
  exactly another row's), a number or date Excel made of typed text, a
  formula, a repeated or unknown row, a renamed column.
- It keeps the walk's rules: a text with a banned phrase can't be accepted; a
  sample stays a sample; an edit without an accept makes a hand-edited draft
  that `generate` never overwrites; an item already reviewed and unchanged is
  left alone (its date is kept).

Everything is planned on a copy of the bill, shown in full, and written once.
"""

from __future__ import annotations

import contextlib
import io
import os
import shutil
import sys
import tempfile
import time
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Callable, Optional, Union, get_args

from openpyxl import load_workbook

from .export import HEADERS, build_workbook
from .generate import refresh_group_sections
from .hashing import sha256_file
from .lint import LANGS, SUMMARY_ID, LintError, group_item_id, lint_bill, load_phrases
from .models import (
    HAND_EDITED,
    Bill,
    Clause,
    GroupEntry,
    GroupId,
    Review,
    Summary,
    Tri,
    bill_path,
    current_sha256,
    is_reviewed,
    load_bill,
    save_bill,
)
from .review import QuitReview, _ask_from_terminal
from .reviewlock import BillLocked, bill_lock
from .selftest import FIXTURE_BILL_ID, Checker, fixture_bill, mark_reviewed

Ask = Callable[[str], str]
Show = Callable[[str], None]
Owner = Union[Clause, Summary, GroupEntry]

# Sheet column for each language, and the only decision that accepts.
LANG_COLUMNS = {"en": "English", "sw": "Kiswahili", "sheng": "Sheng"}
ACCEPT = "accept"
KNOWN_DECISIONS = {"", "accept", "edited", "needs work", "skip"}
GROUP_IDS = set(get_args(GroupId))

# Columns whose cells must be plain text: Excel turns typed "5%" into the
# number 0.05 and "1/7" into a date, and the text saved would not be the
# text shown. Decision and Notes are only read, never saved, so they're free.
TEXT_COLUMNS = ("ID", "Groups / sections", *LANG_COLUMNS.values(), "Legal text", "fingerprint")


class SheetError(Exception):
    """The sheet can't be read safely. Nothing is applied."""


@dataclass
class Row:
    id: str
    decision: str  # lowercased, stripped
    notes: str
    groups: str
    texts: dict[str, str]  # lang -> text, stripped
    legal: str  # the grey Legal text cell, to catch rows scrambled by a sort
    fingerprint: str


@dataclass
class Plan:
    """What the sheet would do, worked out on a copy of the bill."""

    bill: Bill  # the copy, with everything applied
    lines: list[str] = field(default_factory=list)  # one per row that does something
    accepted: list[str] = field(default_factory=list)  # newly reviewed
    edited: list[str] = field(default_factory=list)
    refused: list[str] = field(default_factory=list)
    stale: list[str] = field(default_factory=list)
    unchanged: list[str] = field(default_factory=list)  # Accept on a sample or an item already reviewed
    notes: list[str] = field(default_factory=list)

    @property
    def changes(self) -> bool:
        return bool(self.accepted or self.edited)


# --------------------------------------------------------------------------
# Reading the sheet
# --------------------------------------------------------------------------


def _text(value: object) -> str:
    return "" if value is None else str(value).replace("\r\n", "\n").strip()


def read_sheet(path: Path, bill_id: str) -> list[Row]:
    """The rows of an exported sheet, checked. Raises SheetError on anything
    that could make a row mean something other than what was written."""
    try:
        # Formulas as written (not cached results), so they can be refused.
        book = load_workbook(path)
    except Exception as error:  # openpyxl raises several unrelated types
        raise SheetError(f"can't open {path}: {error}") from error
    if "Review" not in book.sheetnames or "How to use" not in book.sheetnames:
        raise SheetError("not an export sheet: it needs the 'Review' and 'How to use' tabs")

    # The guide tab names the bill it was exported from.
    guide = [_text(row[0]) for row in book["How to use"].iter_rows(max_col=1, values_only=True)]
    named = next((line for line in guide if line.startswith("Bill: ")), "")
    if not named.startswith(f"Bill: {bill_id} ("):
        raise SheetError(f"this sheet is for another bill ({named or 'none named'}), not {bill_id}")

    sheet = book["Review"]
    header = [_text(cell.value) for cell in sheet[1]]
    missing = [name for name in HEADERS if name not in header]
    if missing:
        raise SheetError(f"columns missing or renamed: {missing}")
    col = {name: header.index(name) for name in HEADERS}

    rows: list[Row] = []
    seen: set[str] = set()
    for number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        item_id = _text(cells[col["ID"]].value)
        if not item_id:
            continue
        for name in HEADERS:
            cell = cells[col[name]]
            if cell.data_type == "f":
                raise SheetError(f"{item_id} {name} (row {number}) is a formula; type it as text")
            if name in TEXT_COLUMNS and cell.value is not None and not isinstance(cell.value, str):
                raise SheetError(
                    f"{item_id} {name} (row {number}): Excel stored this as a number or date "
                    f"({cell.value!r}), not the text typed; start the cell with an apostrophe (')"
                )
        if item_id in seen:
            raise SheetError(f"{item_id} appears twice")
        seen.add(item_id)
        value = {name: _text(cells[col[name]].value) for name in HEADERS}
        rows.append(
            Row(
                id=item_id,
                decision=value["Decision"].lower(),
                notes=value["Notes"],
                groups=value["Groups / sections"],
                texts={lang: value[name] for lang, name in LANG_COLUMNS.items()},
                legal=value["Legal text"],
                fingerprint=value["fingerprint"],
            )
        )
    return rows


# --------------------------------------------------------------------------
# Working out the plan
# --------------------------------------------------------------------------


def _owners(bill: Bill) -> dict[str, Owner]:
    owners: dict[str, Owner] = {clause.id: clause for clause in bill.clauses}
    if bill.summary is not None:
        owners[SUMMARY_ID] = bill.summary
    for entry in bill.groups:
        owners[group_item_id(entry.group)] = entry
    return owners


def _tri(owner: Owner) -> Optional[Tri]:
    return owner.explanation if isinstance(owner, Clause) else owner.text


def _parse_groups(item_id: str, raw: str) -> list[str]:
    names = [name.strip().lower() for name in raw.split(",") if name.strip()]
    unknown = [name for name in names if name not in GROUP_IDS]
    if unknown:
        raise SheetError(f"{item_id}: unknown group(s) {unknown}; use {sorted(GROUP_IDS)}")
    return list(dict.fromkeys(names))  # drop repeats, keep order


def _check_rows(owners: dict[str, Owner], rows: list[Row]) -> None:
    """Refuse the whole sheet if any row could mean something other than it
    seems. Runs before anything is applied."""
    unknown = [row.id for row in rows if row.id not in owners]
    if unknown:
        raise SheetError(f"rows for items this bill doesn't have: {unknown}")

    # Each language's original texts, to spot a cell that moved between rows.
    original: dict[str, dict[str, str]] = {lang: {} for lang in LANGS}
    for item_id, owner in owners.items():
        tri = _tri(owner)
        if tri is not None:
            for lang in LANGS:
                original[lang].setdefault(_text(getattr(tri, lang)), item_id)

    for row in rows:
        owner = owners[row.id]
        if isinstance(owner, Clause) and row.legal != _text(owner.text):
            raise SheetError(f"{row.id}: its Legal text cell isn't section {owner.number}'s text. Were rows sorted by one column?")
        tri = _tri(owner)
        if tri is None:
            continue
        for lang in LANGS:
            if not row.texts[lang]:
                raise SheetError(f"{row.id}: the {LANG_COLUMNS[lang]} cell is empty (to drop a text, say so in Notes instead)")
            other = original[lang].get(row.texts[lang])
            if row.texts[lang] != _text(getattr(tri, lang)) and other not in (None, row.id):
                raise SheetError(
                    f"{row.id}: its {LANG_COLUMNS[lang]} is exactly {other}'s text. Was one column sorted on its own?"
                )
        if isinstance(owner, Clause):
            _parse_groups(row.id, row.groups)


def plan(bill: Bill, rows: list[Row], phrases: list[str], today: str) -> Plan:
    """Apply the sheet to a COPY of `bill` and record what happened. Pure:
    `bill` itself is not touched. Raises SheetError before changing anything
    if a row is malformed."""
    work = bill.model_copy(deep=True)
    result = Plan(bill=work)
    owners = _owners(work)

    # 1. Check every row before applying any of them.
    _check_rows(owners, rows)
    for row in rows:
        owner = owners[row.id]
        if isinstance(owner, GroupEntry) and row.groups != ", ".join(owner.clause_ids):
            result.notes.append(f"{row.id}: its sections follow the clause tags; the edit to them here was ignored")
        if row.decision not in KNOWN_DECISIONS:
            result.notes.append(f"{row.id}: decision {row.decision!r} is not one of the dropdown's; NOT accepted")
        if row.notes:
            result.notes.append(f"{row.id} note: {row.notes}")

    # 2. Skip rows whose item moved on in the JSON since the export.
    live: list[Row] = []
    for row in rows:
        if row.fingerprint != (current_sha256(owners[row.id]) or ""):
            result.stale.append(row.id)
            result.lines.append(f"{row.id:<16} SKIPPED: changed in the bill since this sheet was exported")
        else:
            live.append(row)

    sections_before = {entry.group: list(entry.clause_ids) for entry in work.groups}
    actions: dict[str, list[str]] = {row.id: [] for row in live}
    new_texts: dict[str, list[str]] = {row.id: [] for row in live}

    # 3. Text and tag edits.
    for row in live:
        owner = owners[row.id]
        tri = _tri(owner)
        if tri is None:
            tags_moved = isinstance(owner, Clause) and set(_parse_groups(row.id, row.groups)) != set(owner.groups)
            if any(row.texts.values()) or tags_moved:
                result.lines.append(f"{row.id:<16} IGNORED: it has no explanation yet, so nothing here can change")
            continue
        changed = [lang for lang in LANGS if row.texts[lang] != _text(getattr(tri, lang))]
        for lang in changed:
            setattr(tri, lang, row.texts[lang])
            new_texts[row.id].append(f"    {LANG_COLUMNS[lang]}: {row.texts[lang]}")
        actions[row.id] += [f"edited {lang}" for lang in changed]
        if isinstance(owner, Clause):
            groups = _parse_groups(row.id, row.groups)
            if set(groups) != set(owner.groups):
                actions[row.id].append(f"groups {', '.join(owner.groups) or 'none'} -> {', '.join(groups) or 'none'}")
                owner.groups = groups
                changed.append("groups")
        if changed:
            result.edited.append(row.id)
            if owner.review.status != "sample":
                owner.review = Review(status="draft", sha256=HAND_EDITED)

    # 4. Group section lists follow the clause tags (printing what moved).
    with contextlib.redirect_stdout(io.StringIO()) as moved:
        refresh_group_sections(work)
    result.notes += [line.strip() for line in moved.getvalue().splitlines() if line.strip()]
    after = _owners(work)

    # 5. Accepts.
    for row in live:
        if row.decision != ACCEPT:
            continue
        owner = after.get(row.id)
        why = ""
        if owner is None:
            why = "this group was dropped: no section is tagged with it any more"
        elif _tri(owner) is None:
            why = "it has no explanation"
        elif isinstance(owner, GroupEntry) and owner.clause_ids != sections_before.get(owner.group):
            why = (
                f"its sections changed from {', '.join(sections_before.get(owner.group, [])) or 'none'} "
                f"to {', '.join(owner.clause_ids) or 'none'} by tag edits in this sheet; read it again"
            )
        else:
            banned = [f for f in lint_bill(work, phrases) if f.item == row.id and f.kind == "banned"]
            if banned:
                why = "it still has a banned phrase: " + "; ".join(f.describe() for f in banned)
        if why:
            result.refused.append(row.id)
            actions[row.id].append(f"accept REFUSED: {why}")
            continue
        if owner.review.status == "sample":
            result.unchanged.append(row.id)
            actions[row.id].append("accept: stays a sample (only generate replaces a sample)")
        elif is_reviewed(owner):
            result.unchanged.append(row.id)  # same words as already accepted: keep that date
            actions[row.id].append("already reviewed; left as it is")
        else:
            owner.review = Review(status="reviewed", at=today, sha256=current_sha256(owner))
            result.accepted.append(row.id)
            actions[row.id].append("ACCEPT")

    for row in live:
        if actions[row.id]:
            result.lines.append("\n".join([f"{row.id:<16} " + "; ".join(actions[row.id]), *new_texts[row.id]]))
    return result


# --------------------------------------------------------------------------
# `politricks review <bill-id> --from-sheet <file>`
# --------------------------------------------------------------------------


def _save_atomically(bill: Bill, path: Path) -> None:
    """Write beside the bill, then swap it in, so a crash mid-write can never
    leave half a file. The temp name doesn't end in .json, so the site never
    builds a page from it."""
    temp = path.with_name(path.name + ".saving")
    save_bill(bill, temp)
    os.replace(temp, path)


def review_from_sheet(bill_file: Path, sheet_file: Path, phrases: list[str], ask: Ask, show: Show, today: str) -> int:
    """0 = saved (or nothing to do) and lint is clean; 1 otherwise."""
    sha_at_load = sha256_file(bill_file)  # before loading, so the bill read is never newer than this
    bill = load_bill(bill_file)
    try:
        rows = read_sheet(sheet_file, bill.id)
        result = plan(bill, rows, phrases, today)
    except SheetError as error:
        show(f"FAILED: {error}. Nothing was changed.")
        return 1

    show(f"{len(rows)} row(s) read from {sheet_file.name} (sha256 {sha256_file(sheet_file)[:12]}).")
    show("")
    for line in result.lines or ["(no row changes anything)"]:
        show(line)
    if result.notes:
        show("")
        for note in result.notes:
            show(f"  {note}")
    show("")
    show(
        f"accept {len(result.accepted)}, edited {len(result.edited)}, refused {len(result.refused)}, "
        f"skipped as changed since export {len(result.stale)}, left as they were {len(result.unchanged)}"
    )
    if not result.changes:
        show("nothing to save")
        return 0

    try:
        answer = ask("Type yes to save all of this to the bill (anything else quits, saving nothing) > ")
    except (QuitReview, KeyboardInterrupt):
        show("\ninput ended or interrupted: nothing saved")
        return 1
    if answer.strip() != "yes":
        show("not saved")
        return 1
    # run() holds the bill's lock (reviewlock.py), so no politricks command
    # can have written meanwhile. This still catches what the lock can't: a
    # hand edit, or a walk started before the lock existed (2026-10-19).
    if sha256_file(bill_file) != sha_at_load:
        show("NOT saved: the bill file changed while this was open (another review?). Run this again.")
        return 1
    _save_atomically(result.bill, bill_file)
    show(f"saved {bill_file.name}")

    findings = lint_bill(result.bill, phrases)
    remaining = sum(1 for owner in _owners(result.bill).values() if _tri(owner) and not is_reviewed(owner))
    show(f"{remaining} item(s) not reviewed yet; lint findings: {len(findings)}")
    for finding in findings:
        show(finding.line())
    return 0 if not findings else 1


def run(repo_root: Path, bill_id: str, sheet: Path, is_terminal: Callable[[], bool] = sys.stdin.isatty) -> int:
    # The accept must be typed by a person: piped input ("echo yes |") would
    # let a script or an agent answer for him. Refuse before reading anything.
    if not is_terminal():
        print("FAILED: --from-sheet must be run at a terminal by the reviewer, not with input piped in.")
        return 1
    try:
        phrases = load_phrases(repo_root)
        path = bill_path(repo_root, bill_id)
    except (LintError, ValueError) as error:
        print(f"FAILED: {error}")
        return 1
    if not path.exists():
        print(f"FAILED: no bill file at {path}")
        return 1
    # Held from before the bill is read until after it's saved, so a walk or
    # generate run can't write it in between (or overwrite this afterwards).
    try:
        with bill_lock(path):
            return review_from_sheet(path, sheet, phrases, _ask_from_terminal, print, date.today().isoformat())
    except BillLocked as error:
        print(f"FAILED: {error} Nothing was changed.")
        return 1


# --------------------------------------------------------------------------
# Self-test: real export sheets, edited the way a person would, in a temp
# folder. Never the real bill.
# --------------------------------------------------------------------------


def self_test(repo_root: Path) -> bool:
    check = Checker()
    real = bill_path(repo_root, FIXTURE_BILL_ID)
    sha_before = sha256_file(real)
    phrases = load_phrases(repo_root)
    today = time.strftime("%Y-%m-%d")  # from the OS clock, not the code under test

    def cell_edit(sheet_file: Path, item_id: str, **cells: object) -> None:
        """Change cells in one row, by header name, like a person in Excel.
        Underscores in a name stand for spaces ("Groups_/_sections")."""
        book = load_workbook(sheet_file)
        sheet = book["Review"]
        header = [cell.value for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == item_id:
                for name, value in cells.items():
                    row[header.index(name.replace("_", " "))].value = value
        book.save(sheet_file)

    def rename_header(sheet_file: Path, old: str, new: str) -> None:
        book = load_workbook(sheet_file)
        sheet = book["Review"]
        for cell in sheet[1]:
            if cell.value == old:
                cell.value = new
        book.save(sheet_file)

    def setup(tmp: Path, bill: Optional[Bill] = None) -> tuple[Path, Path]:
        bill = bill or fixture_bill(repo_root)
        bill_file, sheet_file = tmp / f"{FIXTURE_BILL_ID}.json", tmp / "sheet.xlsx"
        save_bill(bill, bill_file)
        build_workbook(bill, "2026-01-01").save(sheet_file)
        return bill_file, sheet_file

    def drive(bill_file: Path, sheet_file: Path, answers: list) -> tuple[int, str]:
        lines: list[str] = []
        queue = list(answers)

        def ask(prompt: str) -> str:
            if not queue:
                raise QuitReview("the script ran out")
            answer = queue.pop(0)
            if isinstance(answer, BaseException):
                raise answer
            if callable(answer):  # a side effect at the prompt, then "yes"
                answer()
                return "yes"
            return answer

        code = review_from_sheet(bill_file, sheet_file, phrases, ask, lines.append, today)
        return code, "\n".join(lines)

    def owner(bill_file: Path, item_id: str) -> Optional[Owner]:
        return _owners(load_bill(bill_file)).get(item_id)

    def refused_whole(label: str, edit: Callable[[Path], None], expect: str) -> None:
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", Decision="Accept")
        edit(sheet_file)
        before = sha256_file(bill_file)
        code, shown = drive(bill_file, sheet_file, ["yes"])
        check(
            f"refused whole, nothing written: {label}",
            code == 1 and sha256_file(bill_file) == before and "FAILED" in shown and expect in shown,
            shown,
        )

    BANNED = "The officer cannot vote."

    with tempfile.TemporaryDirectory() as tmp_name:
        tmp = Path(tmp_name)

        # --- A. the main pass ------------------------------------------------
        bill_file, sheet_file = setup(tmp)
        c4_sheng = "Hii section ina-change kitu moja tu."
        cell_edit(sheet_file, "c3", Decision="Accept")
        cell_edit(sheet_file, "c4", Decision="Accept", Sheng=c4_sheng)
        cell_edit(sheet_file, "c1", Decision="Accept")  # a sample
        cell_edit(sheet_file, "group:trader", Decision="accepted")  # a typo on a draft: not an accept
        code, shown = drive(bill_file, sheet_file, ["yes"])
        c3, c4, c1, trader = (owner(bill_file, i) for i in ("c3", "c4", "c1", "group:trader"))
        check("A: accepting c3 unchanged marks it reviewed, dated today", is_reviewed(c3) and c3.review.at == today, str(c3.review))
        check("A: c4's Sheng edit is saved and accepted with the NEW text", c4.explanation.sheng == c4_sheng and is_reviewed(c4))
        check("A: the plan shows c4's new Sheng in full before the prompt", f"Sheng: {c4_sheng}" in shown, shown)
        check("A: Accept on a sample keeps it a sample", c1.review.status == "sample" and "stays a sample" in shown)
        check('A: "accepted" (not a dropdown word) is not an accept', trader.review.status == "draft" and "'accepted' is not one" in shown)

        # --- B. edits without Accept ------------------------------------------
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", Decision="Edited", Kiswahili="Kifungu hiki kinabadilisha jambo moja.")
        cell_edit(sheet_file, "c1", English="A new sample text.")
        drive(bill_file, sheet_file, ["yes"])
        c3, c1 = owner(bill_file, "c3"), owner(bill_file, "c1")
        check("B: an edit marked Edited is a hand-edited draft, not reviewed", c3.review == Review(status="draft", sha256=HAND_EDITED), str(c3.review))
        check("B: an edited sample stays a sample", c1.review.status == "sample" and c1.explanation.en == "A new sample text.", str(c1.review))

        # --- C. nothing is written unless the answer is "yes" -----------------
        for label, answers in (
            ("'y'", ["y"]), ("'YES'", ["YES"]), ("'yes please'", ["yes please"]),
            ("end of input", []), ("Ctrl-C", [KeyboardInterrupt()]),
        ):
            bill_file, sheet_file = setup(tmp)
            cell_edit(sheet_file, "c3", Decision="Accept")
            before = sha256_file(bill_file)
            code, shown = drive(bill_file, sheet_file, answers)
            check(f"C: answering {label} saves nothing and exits 1", sha256_file(bill_file) == before and code == 1, shown[-200:])

        # --- D. a banned phrase can't be accepted; the edit is still kept -----
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", Decision="Accept", English=BANNED)
        cell_edit(sheet_file, "group:trader", Decision="Accept", Sheng=BANNED)
        code, shown = drive(bill_file, sheet_file, ["yes"])
        c3, trader = owner(bill_file, "c3"), owner(bill_file, "group:trader")
        check(
            "D: Accept with a banned phrase is refused; the text is saved as a hand-edited draft",
            c3.explanation.en == BANNED and c3.review.sha256 == HAND_EDITED and "accept REFUSED" in shown,
            shown,
        )
        check("D: the same on a group note", trader.text.sheng == BANNED and not is_reviewed(trader) and "trader" in shown)

        # --- E. a row whose item changed since the export is skipped ---------
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", Decision="Accept", Sheng="Sheng mpya ya sheet.")
        cell_edit(sheet_file, "group:trader", Decision="Accept")
        moved = load_bill(bill_file)
        _owners(moved)["c3"].explanation.en += " Changed in the JSON."
        _owners(moved)["group:trader"].text.en += " Changed in the JSON."
        save_bill(moved, bill_file)
        code, shown = drive(bill_file, sheet_file, ["yes"])
        c3, trader = owner(bill_file, "c3"), owner(bill_file, "group:trader")
        check(
            "E: a clause row changed in the JSON since export is skipped whole (no accept, no edit)",
            not is_reviewed(c3) and c3.explanation.sheng != "Sheng mpya ya sheet." and "c3               SKIPPED" in shown,
            shown,
        )
        check("E: a group row too", not is_reviewed(trader) and "group:trader     SKIPPED" in shown, shown)

        # --- F. group notes and tag edits -------------------------------------
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c4", **{"Groups_/_sections": "employee, trader"})
        cell_edit(sheet_file, "group:trader", Decision="Accept")
        code, shown = drive(bill_file, sheet_file, ["yes"])
        trader, c4 = owner(bill_file, "group:trader"), owner(bill_file, "c4")
        check("F: the tag edit is saved (c4 now trader too)", sorted(c4.groups) == ["employee", "trader"], str(c4.groups))
        check(
            "F: the trader note's sections follow the tags, and its Accept is refused",
            trader.clause_ids == ["c3", "c4"] and not is_reviewed(trader) and "sections changed" in shown,
            shown,
        )
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", **{"Groups_/_sections": "employee"})  # no section left in trader
        cell_edit(sheet_file, "group:trader", Decision="Accept")
        code, shown = drive(bill_file, sheet_file, ["yes"])
        check("F: Accept on a group the tag edits dropped is refused", owner(bill_file, "group:trader") is None and "was dropped" in shown, shown)

        # --- G. an item already reviewed and unchanged is left alone ----------
        reviewed = fixture_bill(repo_root)
        mark_reviewed(_owners(reviewed)["c3"])  # dated 2026-01-01
        bill_file, sheet_file = setup(tmp, reviewed)
        cell_edit(sheet_file, "c3", Decision="Accept")
        before = sha256_file(bill_file)
        code, shown = drive(bill_file, sheet_file, ["yes"])
        check(
            "G: Accept on an unchanged reviewed item keeps its date and asks nothing",
            sha256_file(bill_file) == before and "nothing to save" in shown and "already reviewed" in shown,
            shown,
        )

        # --- H. the bill file changing while the prompt is open --------------
        bill_file, sheet_file = setup(tmp)
        cell_edit(sheet_file, "c3", Decision="Accept")

        def someone_else_saves() -> None:
            meanwhile = load_bill(bill_file)
            _owners(meanwhile)["c4"].explanation.en += " Saved by another walk."
            save_bill(meanwhile, bill_file)

        code, shown = drive(bill_file, sheet_file, [someone_else_saves])
        c3, c4 = owner(bill_file, "c3"), owner(bill_file, "c4")
        check(
            "H: if the bill file changed while the prompt was open, nothing is overwritten",
            code == 1 and not is_reviewed(c3) and c4.explanation.en.endswith("Saved by another walk.") and "NOT saved" in shown,
            shown,
        )
        check("H: no temp file is left beside the bill", not any(path.name.endswith(".saving") for path in tmp.iterdir()))

        # --- I. sheets that may not say what they seem to ---------------------
        refused_whole("an unknown group name", lambda f: cell_edit(f, "c3", **{"Groups_/_sections": "farmer"}), "unknown group")
        refused_whole("an emptied text cell", lambda f: cell_edit(f, "c3", English=""), "is empty")
        refused_whole("a row id the bill doesn't have", lambda f: cell_edit(f, "c4", ID="c999"), "doesn't have")
        refused_whole("a repeated row id", lambda f: cell_edit(f, "c4", ID="c3"), "appears twice")
        refused_whole("a renamed column", lambda f: rename_header(f, "Sheng", "Sheng (new)"), "missing or renamed")
        refused_whole("a number Excel made of typed text", lambda f: cell_edit(f, "c4", Sheng=0.05), "number or date")
        refused_whole("a formula", lambda f: cell_edit(f, "c4", English="=1+1"), "is a formula")

        def swap_english(f: Path) -> None:  # one column sorted on its own
            sheet = load_workbook(f)["Review"]
            values = {row[0].value: row[HEADERS.index("English")].value for row in sheet.iter_rows(min_row=2)}
            cell_edit(f, "c3", English=values["c4"])
            cell_edit(f, "c4", English=values["c3"])

        refused_whole("an English cell that is exactly another row's", swap_english, "exactly")
        refused_whole("a Legal text cell from another row", lambda f: cell_edit(f, "c3", Legal_text="4. Something else."), "Legal text cell")

        bill_file, sheet_file = setup(tmp)
        other = fixture_bill(repo_root)
        other.id = "some-other-bill"
        build_workbook(other, "2026-01-01").save(sheet_file)
        before = sha256_file(bill_file)
        code, shown = drive(bill_file, sheet_file, ["yes"])
        check("I: a sheet exported from another bill is refused", code == 1 and sha256_file(bill_file) == before and "another bill" in shown, shown)

        # --- J. plan() works on a copy ----------------------------------------
        bill = fixture_bill(repo_root)
        bill_file, sheet_file = setup(tmp, bill)
        cell_edit(sheet_file, "c3", Decision="Accept", Sheng="Sheng ya copy.")
        before = bill.model_dump_json()
        plan(bill, read_sheet(sheet_file, bill.id), phrases, today)
        check("J: plan() leaves the bill it was given untouched", bill.model_dump_json() == before)

        # --- K. piped input is refused before anything is read ----------------
        with contextlib.redirect_stdout(io.StringIO()) as said:
            code = run(repo_root, FIXTURE_BILL_ID, sheet_file, is_terminal=lambda: False)
        check("K: run() with input piped in refuses (exit 1)", code == 1 and "at a terminal" in said.getvalue(), said.getvalue())

        # --- L. one writer at a time (reviewlock.py) ---------------------------
        # A temp repo whose bill another writer holds. The sheet is a fresh,
        # unedited export, so even with the lock broken there is nothing to ask:
        # the run can never sit waiting for someone to type "yes".
        root = tmp / "repo"
        (root / "pipeline" / "prompts").mkdir(parents=True)
        shutil.copy(repo_root / "pipeline" / "prompts" / "banned.txt", root / "pipeline" / "prompts")
        held = bill_path(root, FIXTURE_BILL_ID)
        bill = fixture_bill(repo_root)
        save_bill(bill, held)
        clean_sheet = tmp / "clean.xlsx"
        build_workbook(bill, today).save(clean_sheet)
        before = sha256_file(held)
        with bill_lock(held), contextlib.redirect_stdout(io.StringIO()) as said:
            code = run(root, FIXTURE_BILL_ID, clean_sheet, is_terminal=lambda: True)
        check("L: an import while a walk or generate holds the bill refuses (exit 1), bill unchanged",
              code == 1 and "is open in another" in said.getvalue() and sha256_file(held) == before, said.getvalue())
        with contextlib.redirect_stdout(io.StringIO()) as said:
            code = run(root, FIXTURE_BILL_ID, clean_sheet, is_terminal=lambda: True)
        check("L: once the other writer lets go, the same import runs (nothing to save, exit 0)",
              code == 0 and "nothing to save" in said.getvalue(), said.getvalue())

    check("the real bill file is unchanged", sha256_file(real) == sha_before)
    return check.finish()
