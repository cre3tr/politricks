"""Export a bill's texts to an Excel sheet, for reviewing outside the terminal.

`politricks export <bill-id>` writes `exports/<bill-id>-review.xlsx`: one row
per reviewable item (every clause, the summary, each group note) with the
English, Kiswahili and Sheng side by side, the legal text for context, and
the item's state as the site would show it today.

Why it exists: the review walk shows one item at a time in a terminal. A sheet
lets the learner read the whole bill at once, edit in place and note a
decision per row.

What it does NOT do: change the bill. The export only reads the JSON. Nothing
in a returned sheet can mark an item "Reviewed" on its own, because only an
accept in `politricks review` does that (CLAUDE.md). The hidden `fingerprint`
column records each item's `current_sha256` at export time, so whatever
reads the sheet back later can tell which rows changed and whether the JSON
moved on since.
"""

from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .hashing import sha256_file
from .lint import SUMMARY_ID, group_item_id
from .models import (
    HAND_EDITED,
    Bill,
    Clause,
    GroupEntry,
    Review,
    Summary,
    bill_path,
    changed_since_review,
    current_sha256,
    is_hand_edited,
    is_reviewed,
    load_bill,
)
from .selftest import FIXTURE_BILL_ID, FIXTURE_SAMPLE_ID, Checker, fixture_bill, mark_reviewed

# Excel refuses a cell longer than this. A longer text would be cut off
# silently by some readers, so the export stops instead.
EXCEL_CELL_LIMIT = 32_767

# The choices offered in the Decision column (a dropdown). Free text is
# allowed too; these are the ones a later import will understand.
DECISIONS = ("Accept", "Edited", "Needs work", "Skip")

# (header, width, editable). Order is the sheet's column order.
COLUMNS = [
    ("ID", 10, False),
    ("Kind", 9, False),
    ("Heading", 28, False),
    ("State now", 16, False),
    ("Decision", 13, True),
    ("Notes", 30, True),
    ("Groups / sections", 16, True),
    ("English", 60, True),
    ("Kiswahili", 60, True),
    ("Sheng", 60, True),
    ("Legal text", 60, False),
    ("Sheng words flagged", 24, False),
    ("Neutral check", 24, False),
    ("fingerprint", 12, False),
]
HEADERS = [name for name, _, _ in COLUMNS]

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")  # light yellow: yours to change
READ_FILL = PatternFill("solid", fgColor="F2F2F2")  # grey: for context only


def state_label(owner: Clause | Summary | GroupEntry) -> str:
    """The item's state in words, matching the site's badge rules."""
    text = owner.explanation if isinstance(owner, Clause) else owner.text
    if text is None:
        return "No explanation"
    if is_reviewed(owner):
        return "Reviewed"
    if changed_since_review(owner):
        return "Changed since review"
    if owner.review.status == "sample":
        return "Sample"
    if is_hand_edited(owner):
        return "Hand-edited draft"
    return "AI draft"


def _swaps(owner: Clause | Summary | GroupEntry) -> str:
    return "; ".join(f"{swap.word} -> {swap.english}" for swap in owner.sheng_swaps)


def rows(bill: Bill) -> list[list[str]]:
    """One row per item, in the walk's own order: clauses, summary, groups."""
    out: list[list[str]] = []
    for clause in bill.clauses:
        tri = clause.explanation
        neutral = "" if clause.neutral_check is None else (
            ("ok" if clause.neutral_check.ok else "FLAGGED") + (f": {clause.neutral_check.note}" if clause.neutral_check.note else "")
        )
        out.append([
            clause.id, "clause", f"Section {clause.number}: {clause.heading}", state_label(clause), "", "",
            ", ".join(clause.groups),
            tri.en if tri else "", tri.sw if tri else "", tri.sheng if tri else "",
            clause.text, _swaps(clause), neutral, current_sha256(clause) or "",
        ])
    if bill.summary is not None:
        s = bill.summary
        out.append([
            SUMMARY_ID, "summary", "Whole-bill summary", state_label(s), "", "", "",
            s.text.en, s.text.sw, s.text.sheng, "", _swaps(s), "", current_sha256(s) or "",
        ])
    for entry in bill.groups:
        out.append([
            group_item_id(entry.group), "group", f"Note for: {entry.group}", state_label(entry), "", "",
            ", ".join(entry.clause_ids),
            entry.text.en, entry.text.sw, entry.text.sheng, "", _swaps(entry), "", current_sha256(entry) or "",
        ])
    return out


def _put(cell, value: str) -> None:
    """Write a value as plain text. openpyxl stores a string that starts with
    "=" as a formula, so a text like "= 5% of..." would be evaluated by Excel."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def build_workbook(bill: Bill, exported_on: str) -> Workbook:
    data = rows(bill)
    for row in data:
        for header, value in zip(HEADERS, row):
            if len(value) > EXCEL_CELL_LIMIT:
                raise ValueError(f"{row[0]} {header}: {len(value)} characters, over Excel's {EXCEL_CELL_LIMIT}")

    book = Workbook()
    sheet = book.active
    sheet.title = "Review"

    for col, (name, width, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=name)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col)].width = width

    body = Font(name=ARIAL, size=10)
    top_wrap = Alignment(vertical="top", wrap_text=True)
    for r, row in enumerate(data, start=2):
        for col, (value, (_, _, editable)) in enumerate(zip(row, COLUMNS), start=1):
            cell = sheet.cell(row=r, column=col)
            _put(cell, value)
            cell.font = body
            cell.alignment = top_wrap
            cell.fill = EDIT_FILL if editable else READ_FILL

    last = len(data) + 1
    decision_col = get_column_letter(HEADERS.index("Decision") + 1)
    choices = DataValidation(
        type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True, showErrorMessage=False
    )
    choices.add(f"{decision_col}2:{decision_col}{last}")
    sheet.add_data_validation(choices)

    sheet.freeze_panes = "C2"  # ids and kinds stay visible while scrolling
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    sheet.column_dimensions[get_column_letter(HEADERS.index("fingerprint") + 1)].hidden = True

    _guide(book.create_sheet("How to use"), bill, exported_on, len(data))
    return book


def _guide(sheet, bill: Bill, exported_on: str, count: int) -> None:
    lines = [
        ("PoliTricks review sheet", True),
        (f"Bill: {bill.id} ({bill.title})", False),
        (f"Exported: {exported_on}, {count} items", False),
        ("", False),
        ("Yellow cells are yours to change. Grey cells are for context and are not read back.", True),
        ("Decision: pick Accept, Edited, Needs work or Skip from the dropdown (free text is fine too).", False),
        ("Notes: anything you want to say about the row.", False),
        ("English / Kiswahili / Sheng: edit the text in place. Keep every number exactly as in the legal text.", False),
        ("Groups / sections: for a clause, its groups, comma-separated, from: employee, trader, transport, student.", False),
        ("   For a group note it lists that group's sections, which follow the clause tags: change the tags, not this.", False),
        ("State now: what the site shows today (Reviewed, Changed since review, Hand-edited draft, AI draft, Sample).", False),
        ("", False),
        ("Example: c55 | Decision: Edited | Notes: added what the levy is charged on | Sheng: (your new wording)", False),
        ("", False),
        ("Filling in this sheet does not mark anything Reviewed. Only an accept in `politricks review` does.", True),
        ("Don't edit the bill JSON or rerun the export while you hold edits here: send the sheet back first.", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        cell = sheet.cell(row=r, column=1, value=text)
        cell.font = Font(name=ARIAL, size=11 if bold else 10, bold=bold)
    sheet.column_dimensions["A"].width = 110


def export_path(repo_root: Path, bill_id: str) -> Path:
    return repo_root / "exports" / f"{bill_id}-review.xlsx"


def run(repo_root: Path, bill_id: str, out: Optional[Path] = None) -> int:
    bill = load_bill(bill_path(repo_root, bill_id))
    path = out or export_path(repo_root, bill_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(bill, date.today().isoformat()).save(path)

    # Read it back: the count and states come from the saved file, not the plan.
    sheet = load_workbook(path)["Review"]
    states: dict[str, int] = {}
    for (state,) in sheet.iter_rows(min_row=2, min_col=4, max_col=4, values_only=True):
        states[state] = states.get(state, 0) + 1
    print(f"wrote {path}")
    print(f"rows: {sheet.max_row - 1}  " + ", ".join(f"{name}: {n}" for name, n in sorted(states.items())))
    return 0


# --------------------------------------------------------------------------
# Self-test (offline, temp folder, never the real bill)
# --------------------------------------------------------------------------


def self_test(repo_root: Path) -> bool:
    check = Checker()
    real = bill_path(repo_root, FIXTURE_BILL_ID)
    sha_before = sha256_file(real)

    bill = fixture_bill(repo_root)
    c3 = next(c for c in bill.clauses if c.id == "c3")
    c4 = next(c for c in bill.clauses if c.id == "c4")
    # One of each state, on in-memory copies only.
    mark_reviewed(c3)
    mark_reviewed(c4)
    c4.explanation.en += " Changed."  # reviewed, then changed
    trader = bill.groups[0]
    trader.review = Review(status="draft", sha256=HAND_EDITED)
    c5 = next(c for c in bill.clauses if c.id == "c5")
    # A text that looks like a formula must stay text.
    c5_tri = c3.explanation.model_copy()
    c5_tri.en = "=1+1 is not a formula here"
    c5.explanation = c5_tri
    c5.review = Review(status="draft", sha256=HAND_EDITED)

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "sheet.xlsx"
        build_workbook(bill, "2026-01-01").save(path)
        book = load_workbook(path)
        sheet = book["Review"]
        header = [cell.value for cell in sheet[1]]
        got = {row[0]: dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True)}

        check("headers are exactly the column list", header == HEADERS, str(header))
        expected_ids = [c.id for c in bill.clauses] + [SUMMARY_ID] + [group_item_id(g.group) for g in bill.groups]
        check(
            "one row per clause, then the summary, then each group, in that order",
            list(got) == expected_ids,
            f"{len(got)} rows vs {len(expected_ids)}",
        )
        states = {item: got[item]["State now"] for item in (FIXTURE_SAMPLE_ID, "c3", "c4", "c5", "c6", SUMMARY_ID, "group:trader")}
        check(
            "states match the site's rules",
            states == {
                FIXTURE_SAMPLE_ID: "Sample", "c3": "Reviewed", "c4": "Changed since review",
                "c5": "Hand-edited draft", "c6": "No explanation", SUMMARY_ID: "Sample",
                "group:trader": "Hand-edited draft",
            },
            str(states),
        )
        row = got["c3"]
        check(
            "c3's three texts and legal text come back exactly, each in its own column",
            (row["English"], row["Kiswahili"], row["Sheng"], row["Legal text"])
            == (c3.explanation.en, c3.explanation.sw, c3.explanation.sheng, c3.text),
        )
        check(
            "the fingerprint column holds each item's current_sha256",
            row["fingerprint"] == current_sha256(c3) and got["c4"]["fingerprint"] == current_sha256(c4),
        )
        cell = sheet.cell(row=expected_ids.index("c5") + 2, column=HEADERS.index("English") + 1)
        check('a text starting with "=" is stored as text, not a formula', cell.data_type == "s" and cell.value.startswith("=1+1"), cell.data_type)
        check("the fingerprint column is hidden", sheet.column_dimensions[get_column_letter(len(HEADERS))].hidden)
        validations = sheet.data_validations.dataValidation
        check(
            "the Decision column has the dropdown",
            len(validations) == 1 and all(d in validations[0].formula1 for d in DECISIONS),
        )
        check("a 'How to use' sheet is included", "How to use" in book.sheetnames, str(book.sheetnames))

        too_long = bill.model_copy(deep=True)
        too_long.clauses[0].text = "x" * (EXCEL_CELL_LIMIT + 1)
        try:
            build_workbook(too_long, "2026-01-01")
            refused = False
        except ValueError:
            refused = True
        check("a text over Excel's cell limit is refused, not cut", refused)

    check("the real bill file is unchanged", sha256_file(real) == sha_before)
    return check.finish()
